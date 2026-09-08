from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path('/home/ezzanox/Sites/gapura/artifacts/Daftar-Perbaikan-SKPL-DAPL-GAPURA.docx')


ITEMS = [
    ('P0', 'Lengkapi desain antarmuka',
     'Bagian 7.3 Desain Antarmuka hanya berisi judul dan belum menjelaskan layar aplikasi.',
     'Tambahkan rancangan untuk login, dashboard admin, katalog, detail kelas, aktivitas pembelajaran, riwayat, sertifikat, manajemen kelas, dan halaman integrasi. Untuk setiap layar tuliskan aktor, tujuan, komponen, validasi, kondisi kosong, pesan gagal, dan hak akses.',
     'Dokumen desain antarmuka dan prototipe atau tangkapan layar setiap halaman utama.',
     'Semua kebutuhan fungsional dapat ditelusuri ke sedikitnya satu layar dan diuji oleh pengguna bisnis.'),
    ('P0', 'Tetapkan sumber autentikasi dan data pegawai',
     'Dokumen menyebut PRESTIGE, SIMPEG, dan KEDATON secara bergantian sebagai sumber login atau data pengguna.',
     'Tetapkan PRESTIGE sebagai identitas peserta sesuai branch dev. Buat tabel otoritas data: admin lokal, peserta PRESTIGE, data profil, jabatan, status pegawai, dan sertifikat. Hapus penyebutan sistem yang tidak digunakan atau jelaskan perannya secara eksplisit.',
     'Matriks integrasi dan diagram alur login final.',
     'Alur pada dokumen, konfigurasi aplikasi, dan hasil uji login menggunakan sumber data yang sama.'),
    ('P0', 'Selesaikan model peran pengguna',
     'Dokumen masih mendefinisikan Admin, Instructor atau Pemateri, dan Peserta; migration dev menggabungkan Instructor ke Admin tetapi middleware Instructor masih tersedia.',
     'Pilih model final. Jika peran final adalah Admin dan Peserta, pindahkan kewenangan pemateri ke Admin dan hapus referensi serta route Instructor yang tidak dipakai. Jika Instructor tetap diperlukan, batalkan penggabungan role dan buat matriks izin yang jelas.',
     'Matriks role dan permission, migration final, serta daftar route per peran.',
     'Tidak ada role, menu, middleware, atau diagram yang bertentangan.'),
    ('P0', 'Ubah kebutuhan fungsional menjadi dapat diuji',
     'Tabel kebutuhan saat ini dominan berupa daftar menu dan belum memuat aturan bisnis, prasyarat, hasil, maupun kriteria penerimaan.',
     'Beri ID pada setiap kebutuhan, misalnya FR-REV-01. Untuk masing-masing tuliskan aktor, pemicu, prasyarat, alur normal, pengecualian, data yang disimpan, dan acceptance criteria. Tambahkan traceability ke halaman, API, tabel data, dan skenario uji.',
     'Daftar kebutuhan fungsional bernomor dan matriks traceability.',
     'Setiap kebutuhan memiliki minimal satu test case dan hasil uji dapat dinyatakan lulus atau gagal.'),
    ('P0', 'Finalkan kontrak API integrasi',
     'Endpoint, format JSON, dan autentikasi masih bertuliskan contoh. Arah dan metode integrasi KEDATON serta IOP juga belum konsisten dengan API branch dev.',
     'Buat spesifikasi API per sistem: pemilik, arah data, base URL, method, endpoint, header, skema request dan response, validasi, kode error, timeout, retry, idempotency key, limit, dan audit log. Selaraskan dengan endpoint GET yang saat ini disediakan branch dev atau ubah implementasi secara terencana.',
     'Dokumen API OpenAPI atau Swagger dan contoh payload tersanitasi.',
     'Uji integrasi positif, data tidak valid, timeout, duplikasi, dan pemulihan retry berhasil.'),
    ('P0', 'Tetapkan ruang lingkup rilis dan backlog gap',
     'Leaderboard, katalog pembelajaran ASN, riwayat pembelajaran, pembatasan OPD, dan integrasi IOP tercantum di ruang lingkup tetapi belum ditemukan sebagai fitur atau route eksplisit pada branch dev.',
     'Tandai setiap fitur sebagai sudah tersedia, dikembangkan pada rilis ini, atau ditunda. Untuk fitur rilis ini, tambah kebutuhan, desain UI, data model, API, estimasi, dan test case. Jangan memasukkan fitur yang belum disepakati ke baseline penerimaan.',
     'Daftar backlog rilis beserta pemilik, prioritas, dan status.',
     'Ruang lingkup dokumen, backlog, dan branch dev memiliki status fitur yang sama.'),
    ('P0', 'Lengkapi desain data dan aturan integritas',
     'Skema belum menunjukkan dukungan untuk OPD, leaderboard, riwayat, audit integrasi, dan retry. Terdapat typo email_verifiemasd_at, dan password disebut encrypted.',
     'Perbaiki nama kolom menjadi email_verified_at dan ubah keterangan password menjadi hashed. Tambahkan entitas atau atribut yang diperlukan, foreign key, index, unique constraint review per course dan user, batas skor, status sinkronisasi sertifikat, serta audit trail integrasi.',
     'ERD, CDM, data dictionary, dan migration yang selaras.',
     'Migration baru dapat dijalankan dari basis data kosong dan data lama tanpa kehilangan integritas.'),
    ('P0', 'Tuntaskan penanganan ulasan ganda dan nilai',
     'Dokumen meminta satu ulasan per peserta per kelas serta nilai maksimum 100, tetapi desain data belum menyatakan penguncian di tingkat basis data.',
     'Tambahkan unique index pada course_id dan user_id di reviews; validasi aplikasi harus menampilkan pesan yang sesuai. Terapkan validasi rentang skor 0 sampai 100 dan definisikan rumus nilai akhir, pembulatan, dan kondisi lulus.',
     'Migration, request validation, dan test case unit atau feature.',
     'Pengiriman ulasan kedua dan nilai di luar rentang ditolak; perhitungan nilai dapat direproduksi.'),
    ('P1', 'Samakan versi teknologi',
     'Dokumen menyebut Laravel 13, tabel keamanan menyebut Laravel 12, sedangkan composer branch dev menggunakan Laravel 11.',
     'Tentukan versi yang benar untuk rilis ini. Bila upgrade diperlukan, buat rencana upgrade, daftar kompatibilitas paket, pengujian regresi, dan tanggal target. Bila tidak, sesuaikan seluruh dokumen dengan Laravel 11 dan versi PHP yang benar.',
     'Baseline teknologi dan daftar dependensi terverifikasi.',
     'Versi pada dokumen sama dengan composer lock dan lingkungan deployment.'),
    ('P1', 'Selaraskan kontrol keamanan autentikasi',
     'Dokumen mensyaratkan registrasi, aktivasi, reset sandi, dan aturan sandi lokal, sedangkan branch dev menonaktifkan proses tersebut bagi peserta yang dikelola PRESTIGE.',
     'Pisahkan ketentuan untuk admin lokal dan peserta SSO. Nyatakan secara eksplisit pihak yang mengelola password, reset password, aktivasi, lockout, dan pemulihan akun. Hapus kontrol yang tidak relevan agar tidak menjadi kewajiban yang tidak dapat dipenuhi.',
     'Matriks kontrol keamanan per tipe akun.',
     'Audit keamanan menunjukkan tidak ada fitur lokal peserta yang bertentangan dengan SSO.'),
    ('P1', 'Selaraskan CAPTCHA dan pembatasan login',
     'Dokumen mewajibkan Google reCAPTCHA, sedangkan branch dev menggunakan CAPTCHA lokal dan throttle login.',
     'Pilih mekanisme final dan dokumentasikan alasan, konfigurasi, fallback, batas percobaan, periode blokir, serta pemantauan kegagalan. Jika tetap menggunakan CAPTCHA lokal, ganti semua penyebutan Google reCAPTCHA.',
     'Spesifikasi kontrol login dan bukti konfigurasi non-rahasia.',
     'Lima percobaan gagal dan tantangan CAPTCHA diuji sesuai ketentuan final.'),
    ('P1', 'Tegaskan sesi, logging, dan privasi',
     'Konfigurasi dev telah menetapkan idle session 15 menit, tetapi dokumen tidak menjelaskan regenerasi sesi, logout lintas perangkat, retensi log, masking, maupun data yang boleh dicatat.',
     'Dokumentasikan masa sesi, regenerasi setelah login, invalidasi saat perubahan role, kebijakan logout, retensi log, field log yang dilarang, masking PII, akses auditor, dan prosedur penghapusan log.',
     'Kebijakan sesi dan logging serta matriks PII.',
     'Sesi kedaluwarsa sesuai ketentuan dan log tidak memuat password, token, atau PII lengkap.'),
    ('P1', 'Satukan kebijakan unggah file',
     'Batas unggah 10 MB bertentangan dengan ketentuan lain yang membatasi dokumen, foto, dan video secara berbeda.',
     'Tetapkan satu tabel kebijakan file per tipe: ekstensi, MIME type, ukuran maksimum, lokasi storage private, masa retensi, antivirus atau malware scan, penamaan acak, dan hak akses. Terapkan validasi server-side.',
     'Kebijakan unggah file dan test case validasi.',
     'File dengan MIME atau ukuran tidak sesuai ditolak; file tidak dapat dieksekusi dari URL publik.'),
    ('P1', 'Perbaiki diagram dan konsistensi arsitektur',
     'Diagram konteks dan DFD belum konsisten menampilkan semua sistem eksternal yang disebut pada dokumen, khususnya IOP dan KEDATON.',
     'Perbarui diagram konteks, DFD level 0 sampai 2, dan arsitektur layanan berdasarkan matriks integrasi final. Setiap aliran harus memiliki nama data, arah, dan pemilik proses yang jelas.',
     'Diagram sumber yang dapat diedit dan ekspor untuk dokumen.',
     'Setiap aliran pada diagram dapat dipetakan ke kontrak API dan kebutuhan fungsional.'),
    ('P1', 'Lengkapi kebutuhan nonfungsional dan uji beban',
     'Target 20 sampai 40 pengguna bersamaan serta waktu proses kurang dari 5 detik belum memiliki skenario, ukuran data, atau ukuran keberhasilan.',
     'Tetapkan SLO yang terukur: jumlah pengguna, transaksi atau aktivitas per menit, percentile waktu respons, tingkat error, ukuran dataset, durasi uji, dan ambang lulus. Tambahkan rencana kapasitas VPS, database, cache, dan queue.',
     'Dokumen NFR terukur dan laporan uji beban.',
     'Laporan uji menunjukkan target terpenuhi pada skenario yang disetujui.'),
    ('P1', 'Lengkapi rencana pengujian dan serah terima',
     'Jadwal masih menggunakan Bulan X dan belum memiliki PIC, tanggal, kriteria masuk dan keluar, hasil, atau mekanisme perbaikan temuan.',
     'Isi jadwal dengan tanggal aktual, PIC penyedia dan BKPSDM atau Diskominfo, lingkungan uji, data uji, kriteria lulus, bukti yang dikumpulkan, severity defect, dan batas waktu perbaikan. Tambahkan UAT serta berita acara serah terima.',
     'Test plan, test case, UAT sign-off, dan berita acara.',
     'Semua test case P0 dan P1 lulus atau memiliki persetujuan risiko tertulis.'),
    ('P2', 'Rapikan administrasi dan format dokumen',
     'Daftar isi dan sejarah revisi belum terisi. Ada perubahan teks berwarna merah tanpa tracked changes formal, tabel berlanjut tanpa header berulang, serta istilah dan ejaan tidak konsisten.',
     'Isi riwayat revisi, buat daftar isi otomatis, terima atau tolak seluruh perubahan sebelum final, ulangi header tabel, dan normalisasi istilah: Admin, Peserta, PRESTIGE, KEDATON, IOP, autentikasi, peserta, performa, dan Kabupaten.',
     'Dokumen final bersih dengan versi, penulis, tanggal, dan persetujuan.',
     'Tidak ada placeholder, teks revisi berwarna, halaman kosong, atau istilah yang berubah-ubah.'),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color='D9D9D9'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = field
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_labelled_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + ': ')
    r.bold = True
    p.add_run(text)
    return p


def add_item(doc, number, item):
    priority, title, finding, action, output, validation = item
    if title == 'Rapikan administrasi dan format dokumen':
        doc.add_page_break()
    p = doc.add_paragraph(style='Heading 2')
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(f'{number}. {title}')
    details = [
        ('Temuan', finding),
        ('Cara perbaikan', action),
        ('Keluaran yang dikumpulkan', output),
        ('Bukti selesai', validation),
    ]
    for label, text in details:
        add_labelled_paragraph(doc, label, text)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Heading 1'].font.name = 'Aptos Display'
styles['Heading 1'].font.color.rgb = RGBColor(0, 0, 0)
styles['Heading 1'].font.size = Pt(15)
styles['Heading 2'].font.name = 'Aptos'
styles['Heading 2'].font.color.rgb = RGBColor(0, 0, 0)
styles['Heading 2'].font.size = Pt(12)
title_style_ppr = styles['Title']._element.find(qn('w:pPr'))
if title_style_ppr is not None:
    title_border = title_style_ppr.find(qn('w:pBdr'))
    if title_border is not None:
        title_style_ppr.remove(title_border)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('DAFTAR PERBAIKAN SKPL DAN DAPL GAPURA E LEARNING')
title.runs[0].font.color.rgb = RGBColor(0, 0, 0)
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn('w:pBdr'))
if title_border is not None:
    title_ppr.remove(title_border)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run('Acuan penyempurnaan sebelum dokumen dikumpulkan dan dijadikan baseline pengembangan').italic = True

doc.add_paragraph('Dokumen ini merangkum perbaikan yang diperlukan pada SKPL dan DAPL GAPURA E Learning berdasarkan penelaahan dokumen serta keselarasan dengan implementasi pada branch dev. Status dokumen belum siap dikumpulkan sebelum seluruh item prioritas P0 diselesaikan atau memiliki persetujuan risiko tertulis.')

doc.add_paragraph('Urutan penyelesaian', style='Heading 1')
priority_table = doc.add_table(rows=1, cols=3)
priority_table.alignment = WD_TABLE_ALIGNMENT.CENTER
priority_table.style = 'Table Grid'
headers = ['Prioritas', 'Makna', 'Syarat sebelum dikumpulkan']
for cell, text in zip(priority_table.rows[0].cells, headers):
    cell.text = text
    set_cell_shading(cell, '1F4E78')
    set_cell_border(cell)
    for r in cell.paragraphs[0].runs:
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.bold = True
set_repeat_table_header(priority_table.rows[0])
for rowdata in [
    ('P0', 'Pengunci kelayakan dokumen dan implementasi', 'Wajib selesai atau disetujui secara tertulis oleh pemilik risiko'),
    ('P1', 'Penting untuk kesiapan uji dan keamanan operasional', 'Diselesaikan sebelum UAT atau memiliki rencana implementasi disetujui'),
    ('P2', 'Kerapian administrasi dan kualitas penyajian', 'Wajib selesai sebelum versi final diterbitkan'),
]:
    cells = priority_table.add_row().cells
    for cell, text in zip(cells, rowdata):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)

doc.add_paragraph('Daftar perbaikan', style='Heading 1')
for i, item in enumerate(ITEMS, 1):
    add_item(doc, i, item)

doc.add_paragraph('Checklist final sebelum pengumpulan', style='Heading 1')
checklist = doc.add_table(rows=1, cols=3)
checklist.style = 'Table Grid'
checklist.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, text in zip(checklist.rows[0].cells, ['Pemeriksaan', 'Status', 'Paraf atau bukti']):
    cell.text = text
    set_cell_shading(cell, '1F4E78')
    set_cell_border(cell)
    for r in cell.paragraphs[0].runs:
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.bold = True
set_repeat_table_header(checklist.rows[0])
for text in [
    'Semua temuan P0 selesai atau memiliki persetujuan risiko tertulis',
    'Dokumen SKPL dan DAPL konsisten dengan branch dev dan baseline teknologi',
    'Kontrak API, ERD, DFD, desain antarmuka, dan test plan telah diperbarui',
    'Tidak ada placeholder, halaman kosong, revisi berwarna, atau perubahan belum diputuskan',
    'UAT, pengujian integrasi, pengujian fungsi, beban, dan keamanan memiliki bukti hasil',
    'Versi, penulis, tanggal, riwayat revisi, serta persetujuan dokumen telah terisi',
]:
    cells = checklist.add_row().cells
    cells[0].text = text
    cells[1].text = 'Belum / Selesai'
    cells[2].text = ''
    for cell in cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Daftar Perbaikan SKPL dan DAPL GAPURA E Learning | Halaman ')
add_field(footer, 'PAGE')

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
